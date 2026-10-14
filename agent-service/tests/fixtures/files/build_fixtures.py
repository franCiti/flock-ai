"""Builds the binary sample files of this folder (PDF, Word, emails) and the non-UTF-8 text sample.

Run from agent-service/:  uv run python tests/fixtures/files/build_fixtures.py
The generated files are committed; this script documents exactly what they contain.
"""

from email.message import EmailMessage
from pathlib import Path

import docx
from fpdf import FPDF

HERE = Path(__file__).resolve().parent


def build_pdf() -> None:
    pdf = FPDF()
    pdf.set_title("Procedimiento de deploy")
    pdf.set_font("Helvetica", size=12)
    pages = [
        "Deploy a produccion. Los deploys se hacen de lunes a jueves entre las 10 y las 16.",
        "Rollback: si la tasa de errores supera el 2 por ciento, ejecutar el job rollback-prod.",
    ]
    for text in pages:
        pdf.add_page()
        pdf.multi_cell(0, 8, text)
    pdf.output(str(HERE / "sample.pdf"))


def build_scanned_pdf() -> None:
    """A page with only drawings and no text layer, like a scanned document."""
    pdf = FPDF()
    pdf.add_page()
    for i in range(10):
        pdf.rect(20, 20 + i * 15, 170, 8, style="F")
    pdf.output(str(HERE / "scanned.pdf"))


def build_docx() -> None:
    document = docx.Document()
    document.core_properties.title = "Manual de accesos"
    document.add_paragraph("Accesos del equipo", style="Title")
    document.add_paragraph("Cómo pedir cada acceso.")
    document.add_heading("VPN", level=1)
    document.add_paragraph("Se pide con un ticket en el proyecto ACCESOS.")
    document.add_heading("Requisitos", level=2)
    document.add_paragraph("Aprobación del manager directo.", style="List Bullet")
    document.add_paragraph("Notebook de la empresa.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Sistema"
    table.cell(0, 1).text = "Aprobador"
    table.cell(1, 0).text = "VPN"
    table.cell(1, 1).text = "Manager"
    document.save(str(HERE / "sample.docx"))


def build_latin1_text() -> None:
    text = "Configuración del servidor\n\nLa contraseña se renueva cada año en diciembre.\n"
    (HERE / "notes-latin1.txt").write_bytes(text.encode("latin-1"))


def build_emails() -> None:
    """mail.eml: text body and every kind of attachment; mail-html.eml: HTML body only.

    `.msg` samples cannot be built (Outlook format, extract-msg only reads it): the .msg tests are
    skipped until someone adds `tests/fixtures/files/sample.msg` saved from Outlook.
    """
    forwarded = EmailMessage()
    forwarded["From"] = "Lucía Díaz <lucia.diaz@example.com>"
    forwarded["To"] = "Juan Pérez <juan.perez@example.com>"
    forwarded["Subject"] = "Ventana de mantenimiento"
    forwarded["Date"] = "Tue, 10 Feb 2026 09:00:00 -0300"
    forwarded.set_content("El mantenimiento del cluster es el sábado de 2 a 4.")
    forwarded.add_attachment(
        b"Detalle que no se indexa (adjunto de un adjunto).",
        maintype="text",
        subtype="plain",
        filename="detalle.txt",
    )

    mail = EmailMessage()
    mail["From"] = "Ana Gómez <ana.gomez@example.com>"
    mail["To"] = "Juan Pérez <juan.perez@example.com>"
    mail["Cc"] = "Equipo Plataforma <plataforma@example.com>"
    mail["Subject"] = "Accesos para el nuevo ingreso"
    mail["Date"] = "Mon, 16 Feb 2026 10:12:00 -0300"
    mail.set_content("Hola Juan,\n\nTe paso el checklist de accesos y el manual.\n\nSaludos,\nAna")
    mail.add_attachment(
        b"1. VPN\n2. Jenkins\n3. Grafana\n",
        maintype="text",
        subtype="plain",
        filename="checklist.txt",
    )
    mail.add_attachment(
        (HERE / "sample.docx").read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="manual.docx",
    )
    mail.add_attachment(
        (HERE / "image.png").read_bytes(), maintype="image", subtype="png", filename="logo.png"
    )
    mail.add_attachment(
        b"Version 2 del checklist.", maintype="text", subtype="plain", filename="checklist.txt"
    )
    mail.add_attachment(forwarded)
    (HERE / "mail.eml").write_bytes(bytes(mail))

    html = EmailMessage()
    html["From"] = "Monitoreo <alertas@example.com>"
    html["To"] = "Juan Pérez <juan.perez@example.com>"
    html["Subject"] = "Alerta: disco al 90 %"
    html["Date"] = "Wed, 18 Feb 2026 03:00:00 -0300"
    html.set_content(
        "<html><body><h1>Alerta</h1><p>El disco de <b>db-01</b> está al 90 %.</p></body></html>",
        subtype="html",
    )
    (HERE / "mail-html.eml").write_bytes(bytes(html))


if __name__ == "__main__":
    build_pdf()
    build_scanned_pdf()
    build_docx()
    build_latin1_text()
    build_emails()
