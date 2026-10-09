"""Builds the Word and PDF documents of examples/knowledge (fictional personas, plan 8.6).

Run from agent-service/:  uv run python scripts/build_examples.py
The generated files are committed; this script documents exactly what they contain.
"""

from pathlib import Path

import docx
from fpdf import FPDF

EXAMPLES = Path(__file__).resolve().parents[2] / "examples" / "knowledge"


def build_docx(path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    document = docx.Document()
    document.core_properties.title = title
    document.add_paragraph(title, style="Title")
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            if paragraph.startswith("- "):
                document.add_paragraph(paragraph[2:], style="List Bullet")
            else:
                document.add_paragraph(paragraph)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def build_pdf(path: Path, title: str, pages: list[str]) -> None:
    pdf = FPDF()
    pdf.set_title(title)
    pdf.set_font("Helvetica", size=11)
    for text in pages:
        pdf.add_page()
        pdf.multi_cell(0, 7, text)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))


def build_juan() -> None:
    root = EXAMPLES / "juan-plataforma"
    build_docx(
        root / "accesos" / "politica-de-accesos.docx",
        "Política de accesos de Plataforma",
        [
            (
                "Acceso de lectura a la base de producción",
                [
                    "Solo se da acceso de lectura, y por 30 días.",
                    "Se pide con un ticket en el proyecto ACCESOS, tipo 'Acceso DB', indicando el "
                    "motivo. Lo aprueba el líder técnico de Plataforma.",
                    "Nunca se da acceso de escritura a personas: los cambios de datos se hacen con "
                    "scripts revisados.",
                ],
            ),
            (
                "Usuario de Jenkins bloqueado",
                [
                    "Jenkins bloquea el usuario después de 5 intentos fallidos.",
                    "Se desbloquea solo a los 30 minutos; si es urgente, pedirlo en "
                    "#plataforma-soporte.",
                ],
            ),
            (
                "Bajas",
                ["- Las cuentas de quien deja la empresa se dan de baja el mismo día."],
            ),
        ],
    )
    # fpdf2's core fonts only cover latin-1: plain ASCII text keeps the sample simple.
    build_pdf(
        root / "procedimientos" / "runbook-incidentes.pdf",
        "Runbook de incidentes de Plataforma",
        [
            "Runbook de incidentes (documento ficticio).\n\n"
            "Severidad 1: el servicio principal esta caido para todos los usuarios. Se llama al "
            "titular de guardia por telefono y se abre una sala en Meet en menos de 10 minutos.\n\n"
            "Severidad 2: una funcionalidad importante esta degradada. Se avisa en "
            "#plataforma-incidentes y se atiende en menos de 1 hora.",
            "Comunicacion: el responsable del incidente actualiza el estado cada 30 minutos en "
            "#plataforma-incidentes.\n\n"
            "Postmortem: todo incidente de severidad 1 o 2 tiene un postmortem escrito en los "
            "5 dias "
            "habiles siguientes, sin buscar culpables.",
        ],
    )


def build_ana() -> None:
    root = EXAMPLES / "ana-datos"
    build_docx(
        root / "diccionario" / "diccionario-clientes.docx",
        "Diccionario de datos de clientes",
        [
            (
                "Tabla dim_clientes",
                [
                    "- id_cliente: identificador único del cliente en el ERP.",
                    "- segmento: minorista, mayorista o corporativo.",
                    "- fecha_alta: fecha en que el cliente hizo su primera compra.",
                ],
            ),
            (
                "Reglas",
                [
                    "Un cliente pasa a 'inactivo' si no compra durante 12 meses.",
                    "Los datos personales (email, teléfono) no se copian al data warehouse.",
                ],
            ),
        ],
    )
    build_pdf(
        root / "cierres" / "calendario-cierres.pdf",
        "Calendario de cierres mensuales",
        [
            "Calendario de cierres mensuales de ventas (documento ficticio).\n\n"
            "El cierre del mes se hace el tercer dia habil del mes siguiente. Ese dia el "
            "pipeline de "
            "ventas corre dos veces: a las 03:00 y a las 12:00, para incluir las notas de credito "
            "cargadas a ultimo momento.\n\n"
            "Finanzas recibe el reporte de cierre a las 15:00 del mismo dia.",
        ],
    )


if __name__ == "__main__":
    build_juan()
    build_ana()
